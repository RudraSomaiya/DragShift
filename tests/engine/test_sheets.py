"""Spreadsheet converters (PRD §6.1 spreadsheet row). Samples are generated here."""

import datetime as dt
import hashlib
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from dragshift.engine import registry
from dragshift.engine.converters import sheets
from dragshift.engine.worker import JobContext, final_path


def finish(parts: list[Path]) -> list[Path]:
    return [p.rename(final_path(p)) for p in parts]


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture
def book(tmp_path: Path) -> Path:
    """Two sheets: a header row, numbers, a date, text with accents and a formula."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Cost", "Day"])
    ws.append(["Café", 3.5, dt.datetime(2026, 10, 21)])
    ws.append(["Tea", 2, dt.datetime(2026, 10, 22)])
    ws["B4"] = "=B2+B3"
    other = wb.create_sheet("Notes")
    other.append(["Zip", "007"])  # text that looks like a number stays text
    path = tmp_path / "plan.xlsx"
    wb.save(path)
    return path


def read_csv(path: Path, sep: str = ",") -> list[list[str]]:
    import csv

    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f, delimiter=sep))


def test_xlsx_to_csv_writes_one_file_per_sheet(book: Path) -> None:
    before = sha(book)
    outs = finish(sheets.to_csv([book], {}, JobContext("t")))
    assert [p.name for p in outs] == ["plan Budget.csv", "plan Notes.csv"]
    rows = read_csv(outs[0])
    assert rows[:3] == [
        ["Item", "Cost", "Day"],
        ["Café", "3.5", "2026-10-21"],
        ["Tea", "2", "2026-10-22"],
    ]
    assert read_csv(outs[1]) == [["Zip", "007"]]
    assert outs[0].read_bytes().startswith(b"\xef\xbb\xbf")  # BOM: Excel reads it as UTF-8
    assert sha(book) == before  # the source is never touched (rule 4)


def test_one_sheet_keeps_the_plain_name(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.active.append(["a", 1])
    wb.save(src := tmp_path / "one.xlsx")
    (out,) = finish(sheets.to_tsv([src], {}, JobContext("t")))
    assert out.name == "one.tsv"
    assert read_csv(out, "\t") == [["a", "1"]]


@pytest.mark.parametrize(("func", "engine"), [(sheets.to_ods, "odf"), (sheets.to_xlsx, "openpyxl")])
def test_workbooks_keep_every_sheet(book: Path, tmp_path: Path, func, engine: str) -> None:
    src = book
    if func is sheets.to_xlsx:  # XLSX is hidden for .xlsx: go through ODS first
        (src,) = finish(sheets.to_ods([book], {}, JobContext("t")))
    (out,) = finish(func([src], {}, JobContext("t")))
    frames = pd.read_excel(out, sheet_name=None, header=None, engine=engine, dtype=object)
    assert list(frames) == ["Budget", "Notes"]
    assert frames["Budget"].iloc[1, 0] == "Café"
    assert frames["Budget"].iloc[1, 1] == 3.5
    assert frames["Notes"].iloc[0, 1] == "007"


def test_european_csv_with_semicolons_cp1252_and_ragged_rows(tmp_path: Path) -> None:
    src = tmp_path / "export.csv"
    src.write_bytes("Name;Prix;Note\r\nCrème;12;\r\nThé;7;chaud;extra\r\n".encode("cp1252"))
    (out,) = finish(sheets.to_xlsx([src], {}, JobContext("t")))
    ws = openpyxl.load_workbook(out).active
    assert ws.title == "export"
    assert [c.value for c in ws[2]][:2] == ["Crème", 12]  # a real number, not text
    assert ws["D3"].value == "extra"  # the longer row is not lost


def test_csv_numbers(tmp_path: Path) -> None:
    assert sheets.number("42") == 42
    assert sheets.number("-3.5") == -3.5
    assert sheets.number("1e3") == 1000.0
    for text in ("007", "1 234", "12,5", "1234567890123456", "abc", ""):
        assert sheets.number(text) == text  # Excel would mangle these


def test_sheet_names() -> None:
    assert sheets.sheet_name("Q1/Q2: [draft]") == "Q1_Q2_ _draft_"
    assert len(sheets.sheet_name("x" * 40)) == 31
    assert sheets.sheet_name("") == "Sheet1"


def test_text_cells() -> None:
    assert sheets.text_cell(float("nan")) == ""
    assert sheets.text_cell(3.0) == 3
    assert sheets.text_cell(dt.datetime(2026, 1, 2)) == "2026-01-02"
    assert sheets.text_cell(dt.datetime(2026, 1, 2, 9, 30)) == dt.datetime(2026, 1, 2, 9, 30)
    assert sheets.text_cell(pd.Timestamp("2026-01-02")) == "2026-01-02"
    assert sheets.text_cell(pd.NaT) == ""


def test_petals() -> None:
    def labels(name: str) -> list[str]:
        return [o.label for o in registry.ops_for([Path(name)], "convert")]

    assert labels("a.xlsx") == ["CSV", "TSV", "ODS"]
    assert labels("a.xls") == ["XLSX", "CSV", "TSV", "ODS"]
    assert labels("a.csv") == ["XLSX", "TSV", "ODS"]
