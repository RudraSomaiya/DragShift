"""PDF converters (PRD §6.1 PDF row): DOCX, PNG/JPG pages, TXT, MD, PPTX, and odd inputs.

Every sample PDF is built here with PyMuPDF; nothing binary is committed. Converters run through
``run_job`` with a plain JobContext, except for the end-to-end tests at the bottom, which use the
real worker pool.
"""

import hashlib
import logging
import subprocess
import sys
from collections.abc import Iterator
from pathlib import Path

import docx
import pymupdf
import pytest
from PIL import Image
from pptx import Presentation

from dragshift.engine import registry
from dragshift.engine.converters import pdf
from dragshift.engine.jobs import JobManager
from dragshift.engine.worker import JobContext, final_path, run_job

MODULE = "dragshift.engine.converters.pdf"
A4 = (595, 842)  # points
LANDSCAPE = (842, 595)


# --- helpers ---------------------------------------------------------------------------------
def make_pdf(path: Path, pages: int = 3, sizes: list[tuple[int, int]] | None = None) -> Path:
    doc = pymupdf.open()
    for i in range(pages):
        w, h = sizes[i] if sizes else A4
        page = doc.new_page(width=w, height=h)
        page.insert_text((72, 72), f"Heading {i + 1}", fontsize=24)
        page.insert_text((72, 120), f"Body text on page {i + 1}.", fontsize=11)
    doc.save(path)
    doc.close()
    return path


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def run(func: str, src: Path, params: dict | None = None, ctx: JobContext | None = None) -> dict:
    """Run an Op the way the pool does, check the source is untouched, finish the outputs."""
    before = sha(src)
    result = run_job(f"{MODULE}:{func}", [[src]], params or {}, ctx or JobContext("t"))
    assert sha(src) == before  # CLAUDE.md rule 4
    result["outputs"] = [part.rename(final_path(part)) for part in result["outputs"]]
    return result


def converted(func: str, src: Path, params: dict | None = None) -> list[Path]:
    result = run(func, src, params)
    assert result["ok"], result["message"]
    return result["outputs"]


def tree(folder: Path) -> list[str]:
    return sorted(str(p.relative_to(folder)) for p in folder.rglob("*"))


# --- PNG / JPG pages (FR-5.1) ----------------------------------------------------------------
def test_pages_go_into_a_pages_folder_at_the_dpi(tmp_path: Path) -> None:
    src = make_pdf(tmp_path / "doc.pdf")
    outputs = converted("to_png", src, {"dpi": 72})
    assert [p.name for p in outputs] == ["doc 001.png", "doc 002.png", "doc 003.png"]
    assert {p.parent.name for p in outputs} == {"doc pages"}
    for p in outputs:
        with Image.open(p) as im:
            assert (im.format, im.size, im.mode) == ("PNG", A4, "RGB")


def test_dpi_scales_the_pixel_size(tmp_path: Path) -> None:
    src = make_pdf(tmp_path / "doc.pdf", pages=1)
    (out,) = converted("to_png", src, {"dpi": 144})
    with Image.open(out) as im:
        assert im.size == (A4[0] * 2, A4[1] * 2)


def test_default_dpi_is_300(tmp_path: Path) -> None:
    src = make_pdf(tmp_path / "doc.pdf", pages=1, sizes=[(72, 36)])
    (out,) = converted("to_jpg", src)
    with Image.open(out) as im:
        assert im.size == (300, 150)


@pytest.mark.parametrize("dpi", [0, -5, 10])
def test_dpi_has_a_floor(tmp_path: Path, dpi: int) -> None:
    src = make_pdf(tmp_path / "doc.pdf", pages=1, sizes=[(72, 72)])
    (out,) = converted("to_png", src, {"dpi": dpi})
    with Image.open(out) as im:
        assert im.size == (pdf.MIN_DPI, pdf.MIN_DPI)


def test_one_page_goes_next_to_the_source(tmp_path: Path) -> None:
    src = make_pdf(tmp_path / "doc.pdf", pages=1)
    (out,) = converted("to_jpg", src, {"dpi": 72})
    assert out == tmp_path / "doc.jpg"
    with Image.open(out) as im:
        assert im.format == "JPEG"


def test_jpg_quality_comes_from_params(tmp_path: Path) -> None:
    src = make_pdf(tmp_path / "doc.pdf", pages=1)
    (low,) = converted("to_jpg", src, {"dpi": 150, "quality": 10})
    (high,) = converted("to_jpg", src, {"dpi": 150, "quality": 95})
    assert low.stat().st_size < high.stat().st_size


def test_an_existing_pages_folder_is_left_alone(tmp_path: Path) -> None:
    src = make_pdf(tmp_path / "doc.pdf", pages=2)
    old = tmp_path / "doc pages"
    old.mkdir()
    (old / "keep.txt").write_text("mine", encoding="utf-8")
    outputs = converted("to_png", src, {"dpi": 36})
    assert {p.parent.name for p in outputs} == {"doc pages (2)"}
    assert tree(old) == ["keep.txt"]


def test_a_dotted_name_keeps_its_whole_stem(tmp_path: Path) -> None:
    src = make_pdf(tmp_path / "v1.2 report.pdf", pages=2)
    outputs = converted("to_png", src, {"dpi": 36})
    assert outputs[0].relative_to(tmp_path) == Path("v1.2 report pages/v1.2 report 001.png")


def test_huge_pages_are_capped(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(pdf, "MAX_SIDE", 400)
    src = make_pdf(tmp_path / "poster.pdf", pages=1, sizes=[(1000, 500)])
    (out,) = converted("to_png", src, {"dpi": 300})
    with Image.open(out) as im:
        assert max(im.size) <= 400


# --- text formats ----------------------------------------------------------------------------
def test_txt_separates_pages_with_a_form_feed(tmp_path: Path) -> None:
    src = make_pdf(tmp_path / "doc.pdf")
    (out,) = converted("to_txt", src)
    assert out == tmp_path / "doc.txt"
    pages = out.read_text(encoding="utf-8").split("\f")
    assert len(pages) == 3
    assert "Body text on page 2." in pages[1]


def test_md_keeps_the_text_of_every_page(tmp_path: Path) -> None:
    src = make_pdf(tmp_path / "doc.pdf")
    (out,) = converted("to_md", src)
    assert out == tmp_path / "doc.md"
    text = out.read_text(encoding="utf-8")
    for i in (1, 2, 3):
        assert f"Body text on page {i}." in text
    assert text.index("page 1") < text.index("page 2") < text.index("page 3")


def test_docx_holds_the_text_of_every_page(tmp_path: Path) -> None:
    src = make_pdf(tmp_path / "doc.pdf")
    (out,) = converted("to_docx", src)
    assert out == tmp_path / "doc.docx"
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    for i in (1, 2, 3):
        assert f"Body text on page {i}." in text


def fail_on_pages(monkeypatch: pytest.MonkeyPatch, *ids: int) -> None:
    """Make pdf2docx fail to parse these page indexes."""
    from pdf2docx.page.Page import Page as DocxPage

    real = DocxPage.parse

    def parse(self, **settings):
        if self.id in ids:
            raise RuntimeError(f"cannot parse page {self.id}")
        return real(self, **settings)

    monkeypatch.setattr(DocxPage, "parse", parse)


def test_docx_skips_a_page_pdf2docx_cannot_read(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fail_on_pages(monkeypatch, 1)
    src = make_pdf(tmp_path / "doc.pdf")
    (out,) = converted("to_docx", src)
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "page 1." in text and "page 3." in text
    assert "page 2." not in text


def test_docx_fails_when_no_page_can_be_read(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fail_on_pages(monkeypatch, 0, 1, 2)
    src = make_pdf(tmp_path / "doc.pdf")
    result = run("to_docx", src)
    assert not result["ok"]
    assert "could not read any page" in result["message"]
    assert tree(tmp_path) == ["doc.pdf"]


def test_importing_pdf2docx_leaves_the_root_logger_alone() -> None:
    """pdf2docx runs logging.basicConfig(INFO) at import; in a fresh process (this one may have
    imported it already) the worker's root logger must stay unconfigured."""
    code = (
        "import logging\n"
        "from dragshift.engine.converters.pdf import _import_pdf2docx\n"
        "_import_pdf2docx()\n"
        "root = logging.getLogger()\n"
        "print(len(root.handlers), root.level)\n"
    )
    out = subprocess.run(
        [sys.executable, "-c", code], capture_output=True, text=True, check=True, timeout=120
    )
    assert out.stdout.split() == ["0", str(logging.WARNING)]  # no handler, Python's default
    assert "fitz" not in out.stderr


# --- PPTX ------------------------------------------------------------------------------------
def test_pptx_has_one_picture_slide_per_page_sized_like_the_page(tmp_path: Path) -> None:
    src = make_pdf(tmp_path / "doc.pdf")
    (out,) = converted("to_pptx", src)
    prs = Presentation(str(out))
    assert (prs.slide_width, prs.slide_height) == (A4[0] * 12700, A4[1] * 12700)
    assert len(prs.slides) == 3
    for slide in prs.slides:
        (pic,) = slide.shapes
        assert pic.shape_type == 13  # MSO_SHAPE_TYPE.PICTURE
        assert (pic.width, pic.height) == (prs.slide_width, prs.slide_height)


def test_pptx_fits_and_centres_a_page_of_another_size(tmp_path: Path) -> None:
    src = make_pdf(tmp_path / "doc.pdf", pages=2, sizes=[A4, LANDSCAPE])
    (out,) = converted("to_pptx", src)
    prs = Presentation(str(out))
    (pic,) = prs.slides[1].shapes
    assert pic.width == prs.slide_width
    assert pic.height < prs.slide_height
    assert abs(pic.top * 2 + pic.height - prs.slide_height) <= 2  # centred


def test_pptx_images_are_capped_at_150_dpi(tmp_path: Path) -> None:
    src = make_pdf(tmp_path / "doc.pdf", pages=1, sizes=[(72, 72)])
    (out,) = converted("to_pptx", src, {"dpi": 600})
    (pic,) = Presentation(str(out)).slides[0].shapes
    assert pic.image.size == (150, 150)


# --- bad inputs and cancel -------------------------------------------------------------------
@pytest.mark.parametrize("func", ["to_docx", "to_png", "to_txt", "to_md", "to_pptx"])
def test_a_password_protected_pdf_fails_cleanly(tmp_path: Path, func: str) -> None:
    src = tmp_path / "locked.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(src, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="u", owner_pw="o")
    doc.close()
    result = run(func, src)
    assert not result["ok"]
    assert "password-protected" in result["message"]
    assert tree(tmp_path) == ["locked.pdf"]


def test_a_broken_pdf_fails_cleanly(tmp_path: Path) -> None:
    src = tmp_path / "broken.pdf"
    src.write_bytes(b"this is not a pdf")
    result = run("to_txt", src)
    assert not result["ok"]
    assert "not a readable PDF" in result["message"]
    assert tree(tmp_path) == ["broken.pdf"]


class CancelAfter:
    """A cancel flag that turns on after ``n`` checks."""

    def __init__(self, n: int) -> None:
        self.n = n

    def __contains__(self, _job: object) -> bool:
        self.n -= 1
        return self.n < 0


@pytest.mark.parametrize("func", ["to_docx", "to_png", "to_txt", "to_md", "to_pptx"])
def test_cancel_part_way_leaves_nothing(tmp_path: Path, func: str) -> None:
    src = make_pdf(tmp_path / "doc.pdf")
    result = run(func, src, {"dpi": 36}, JobContext("t", None, CancelAfter(3)))
    assert result["message"] == "Cancelled"
    assert tree(tmp_path) == ["doc.pdf"]  # no .part files and no empty pages folder


# --- petals (§6.1 order) ---------------------------------------------------------------------
def labels(*names: str) -> list[str]:
    return [o.label for o in registry.ops_for([Path(n) for n in names], "convert")]


def test_petals_follow_the_prd_order() -> None:
    assert labels("a.pdf") == ["DOCX", "PNG", "JPG", "TXT", "MD", "PPTX"]
    assert labels("a.PDF", "b.pdf") == ["DOCX", "PNG", "JPG", "TXT", "MD", "PPTX"]
    assert labels("a.pdf", "b.png") == []  # mixed kinds


# --- through the real worker pool ------------------------------------------------------------
@pytest.fixture
def jm(qtbot) -> Iterator[JobManager]:
    manager = JobManager()
    yield manager
    manager.shutdown()


@pytest.mark.parametrize("ext", ["docx", "png", "md", "pptx"])
def test_converts_in_a_worker_process(qtbot, jm: JobManager, tmp_path: Path, ext: str) -> None:
    src = make_pdf(tmp_path / "doc.pdf")
    with qtbot.waitSignal(jm.finished, timeout=120_000) as blocker:
        jm.submit(f"convert.pdf.{ext}", [src], {"dpi": 72})
    result = blocker.args[1]
    assert result.ok, result.message
    assert result.outputs and all(p.exists() and p.suffix == f".{ext}" for p in result.outputs)
