"""PDF converters (PRD §6.1 PDF row).

Targets in default petal order: DOCX PNG JPG TXT MD PPTX. PNG and JPG render every page at the
Settings > Output DPI (``params["dpi"]``). A one-page PDF gives ``<stem>.png`` next to the source;
more pages go into a new ``<stem> pages`` folder as ``<stem> 001.png``, … (FR-5.1). PPTX holds
one page image per slide. No OCR: a scanned page gives no text.

PyMuPDF, pdf2docx, pymupdf4llm and python-pptx are imported inside the functions: the UI process
imports this module to register Ops (CLAUDE.md rule 2).
"""

import io
import logging
import warnings
from collections.abc import Iterator
from contextlib import contextmanager
from pathlib import Path
from typing import TYPE_CHECKING

from dragshift.engine.converters.images import QUALITY
from dragshift.engine.kinds import Kind
from dragshift.engine.registry import op
from dragshift.engine.worker import new_folder

if TYPE_CHECKING:
    from pymupdf import Document, Page, Pixmap

    from dragshift.engine.worker import JobContext

PDF = frozenset({Kind.PDF})
DPI = 300  # default of Settings > Output > PDF raster DPI (FR-8.1)
MIN_DPI = 36  # half of screen size; anything lower is a mistake in config.json
PPTX_MAX_DPI = 150  # slide images: sharp on screen without a deck of hundreds of MB
MAX_SIDE = 16384  # rendered pages never get larger than this, so a poster cannot exhaust memory
PPTX_MIN_SIDE, PPTX_MAX_SIDE = 914400, 51206400  # PowerPoint's slide size limits: 1 in to 56 in
EMU_PER_PT = 12700


class Cancelled(Exception):
    """The job was cancelled part-way; run_job discards what was reserved."""


# --- reading ---------------------------------------------------------------------------------
@contextmanager
def _open(src: Path) -> Iterator["Document"]:
    import pymupdf

    try:
        doc = pymupdf.open(src, filetype="pdf")
    except pymupdf.FileDataError as e:
        raise ValueError("not a readable PDF") from e
    try:
        if doc.needs_pass:
            raise ValueError("the PDF is password-protected")
        if doc.page_count == 0:
            raise ValueError("the PDF has no pages")
        yield doc
    finally:
        doc.close()


def _pages(doc: "Document", ctx: "JobContext") -> Iterator[int]:
    """Page numbers, reporting progress before each and stopping on cancel."""
    n = doc.page_count
    for i in range(n):
        if ctx.cancelled():
            raise Cancelled
        ctx.progress(i / n, f"Page {i + 1} of {n}")
        yield i


def _dpi(params: dict) -> int:
    """``params["dpi"]``, at least MIN_DPI: MuPDF reads 0 as 72 and a negative DPI as a tiny
    image, and config.json only checks that the value is a whole number."""
    return max(MIN_DPI, int(params.get("dpi", DPI)))


def _render(page: "Page", dpi: int) -> "Pixmap":
    """``page`` as an opaque RGB pixmap at ``dpi``, capped at MAX_SIDE pixels a side."""
    longest_pt = max(page.rect.width, page.rect.height)
    dpi = min(dpi, int(MAX_SIDE * 72 / longest_pt))  # PyMuPDF takes whole DPI only
    return page.get_pixmap(dpi=dpi, alpha=False)


# --- writing ---------------------------------------------------------------------------------
def _rasterise(paths: list[Path], params: dict, ctx: "JobContext", ext: str) -> list[Path]:
    quality = int(params.get("quality", QUALITY))
    outputs = []
    for src in paths:
        with _open(src) as doc:
            if doc.page_count == 1:
                folder, labels = None, [None]
            else:
                folder = new_folder(src.with_name(f"{src.stem} pages"), ctx)
                width = max(3, len(str(doc.page_count)))
                labels = [f"{i + 1:0{width}d}" for i in range(doc.page_count)]
            for i in _pages(doc, ctx):
                dst = ctx.output_path(src, ext=ext, label=labels[i], folder=folder)
                pix = _render(doc[i], _dpi(params))
                if ext == "jpg":
                    pix.save(str(dst), output="jpg", jpg_quality=quality)
                else:
                    pix.save(str(dst), output="png")
                outputs.append(dst)
    return outputs


def _write_text(dst: Path, text: str) -> None:
    dst.write_text(text, encoding="utf-8")


def _import_pdf2docx() -> type:
    """pdf2docx's ``Converter``, imported without its side effects on this worker process: the
    import runs ``logging.basicConfig(level=INFO)``, which would print every later job's INFO
    logs to the console, and imports the deprecated ``fitz`` name (a FutureWarning)."""
    root = logging.getLogger()
    handlers, level = root.handlers[:], root.level
    try:
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", "The `fitz` API is deprecated", FutureWarning)
            from pdf2docx import Converter
    finally:
        root.handlers[:] = handlers
        root.setLevel(level)
    return Converter


# --- Ops, in default petal order (§6.1) ------------------------------------------------------
@op(id="convert.pdf.docx", wheel="convert", label="DOCX", kinds=PDF)
def to_docx(paths: list[Path], params: dict, ctx: "JobContext") -> list[Path]:
    """pdf2docx, driven step by step so each page reports progress and can be cancelled."""
    converter = _import_pdf2docx()
    outputs = []
    for src in paths:
        with _open(src):  # readable, unlocked and not empty, with a clear message if not
            pass
        cv = converter(str(src))
        try:
            settings = cv.default_settings
            cv.load_pages().parse_document(**settings)
            n = len(cv.pages)
            skipped = 0
            for i, page in enumerate(cv.pages):
                if ctx.cancelled():
                    raise Cancelled
                ctx.progress(i / (n + 1), f"Page {i + 1} of {n}")
                try:
                    page.parse(**settings)
                except Exception:  # noqa: BLE001 - as pdf2docx's ignore_page_error, skip the page
                    skipped += 1  # make_docx leaves unparsed pages out
            if skipped == n:
                raise ValueError("pdf2docx could not read any page of this PDF")
            dst = ctx.output_path(src, ext="docx")
            ctx.progress(n / (n + 1), "Writing DOCX")
            cv.make_docx(str(dst), **settings)
        finally:
            cv.close()
        outputs.append(dst)
    return outputs


@op(id="convert.pdf.png", wheel="convert", label="PNG", kinds=PDF)
def to_png(paths: list[Path], params: dict, ctx: "JobContext") -> list[Path]:
    return _rasterise(paths, params, ctx, "png")


@op(id="convert.pdf.jpg", wheel="convert", label="JPG", kinds=PDF)
def to_jpg(paths: list[Path], params: dict, ctx: "JobContext") -> list[Path]:
    return _rasterise(paths, params, ctx, "jpg")


@op(id="convert.pdf.txt", wheel="convert", label="TXT", kinds=PDF)
def to_txt(paths: list[Path], params: dict, ctx: "JobContext") -> list[Path]:
    """The text layer, pages separated by a form feed (``\\f``)."""
    outputs = []
    for src in paths:
        with _open(src) as doc:
            text = "\f".join(doc[i].get_text() for i in _pages(doc, ctx))
        dst = ctx.output_path(src, ext="txt")
        _write_text(dst, text)
        outputs.append(dst)
    return outputs


@op(id="convert.pdf.md", wheel="convert", label="MD", kinds=PDF)
def to_md(paths: list[Path], params: dict, ctx: "JobContext") -> list[Path]:
    """pymupdf4llm, one page per call so each page reports progress and can be cancelled."""
    import pymupdf4llm

    outputs = []
    for src in paths:
        with _open(src) as doc:
            text = "\n".join(
                pymupdf4llm.to_markdown(doc, pages=[i], use_ocr=False) for i in _pages(doc, ctx)
            )
        dst = ctx.output_path(src, ext="md")
        _write_text(dst, text)
        outputs.append(dst)
    return outputs


@op(id="convert.pdf.pptx", wheel="convert", label="PPTX", kinds=PDF)
def to_pptx(paths: list[Path], params: dict, ctx: "JobContext") -> list[Path]:
    """One slide per page, sized like the first page; other page sizes are fitted and centred."""
    from pptx import Presentation

    dpi = min(_dpi(params), PPTX_MAX_DPI)
    outputs = []
    for src in paths:
        prs = Presentation()
        blank = prs.slide_layouts[6]
        with _open(src) as doc:
            first = doc[0].rect
            scale = min(1.0, PPTX_MAX_SIDE / (max(first.width, first.height) * EMU_PER_PT))
            prs.slide_width = max(PPTX_MIN_SIDE, int(first.width * EMU_PER_PT * scale))
            prs.slide_height = max(PPTX_MIN_SIDE, int(first.height * EMU_PER_PT * scale))
            for i in _pages(doc, ctx):
                rect = doc[i].rect
                fit = min(prs.slide_width / rect.width, prs.slide_height / rect.height)
                w, h = round(rect.width * fit), round(rect.height * fit)
                left, top = (prs.slide_width - w) // 2, (prs.slide_height - h) // 2
                image = io.BytesIO(_render(doc[i], dpi).tobytes("png"))
                prs.slides.add_slide(blank).shapes.add_picture(image, left, top, w, h)
        dst = ctx.output_path(src, ext="pptx")
        prs.save(str(dst))
        outputs.append(dst)
    return outputs
